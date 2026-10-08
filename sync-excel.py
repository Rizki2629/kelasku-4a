#!/usr/bin/env python3
"""Parse Daftar Hadir dari ADMINISTRASI KELAS.xlsm -> data.json untuk aplikasi KelasKu 4A."""
import openpyxl, json, sys, datetime

XLSM = sys.argv[1] if len(sys.argv) > 1 else None
OUT = "/home/hatch/workspace/kelas-app/data.json"

BULAN = {"Januari":1,"Februari":2,"Maret":3,"April":4,"Mei":5,"Juni":6,
         "Juli":7,"Agustus":8,"September":9,"Oktober":10,"November":11,"Desember":12}
MAP = {".":"H","s":"S","i":"I","a":"A"}

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    # cari sheet Daftar Hadir (case-insensitive)
    ws = None
    for n in wb.sheetnames:
        if n.lower() == "daftar hadir":
            ws = wb[n]; break
    if ws is None:
        raise SystemExit("Sheet 'Daftar Hadir' tidak ditemukan")
    c4 = str(ws["C4"].value or "").strip()          # "Oktober 2026"
    parts = c4.split()
    bulan = BULAN.get(parts[0], 0); tahun = int(parts[1]) if len(parts)>1 else 0
    if not bulan or not tahun:
        raise SystemExit(f"C4 tidak valid: {c4!r}")
    # baris 7: tanggal di kolom D(4)..AH(34)
    tgl_col = {}
    for c in range(4, 35):
        v = ws.cell(row=7, column=c).value
        if isinstance(v, int) and 1 <= v <= 31:
            tgl_col[c] = v
    att = {}
    n_data = 0; n_siswa = 0
    for r in range(8, 40):
        nama = ws.cell(row=r, column=3).value
        if not nama: continue
        nama = str(nama).strip()
        n_siswa += 1
        for c, tgl in tgl_col.items():
            v = ws.cell(row=r, column=c).value
            if v is None: continue
            v = str(v).strip().lower()
            if v in MAP:
                key = f"{tahun:04d}-{bulan:02d}-{tgl:02d}"
                att.setdefault(key, {})[nama] = MAP[v]
                n_data += 1
    return {"bulan": c4, "attendance": att, "n_data": n_data, "n_siswa": n_siswa}



def parse_catatan(wb):
    """Parse sheet Catatan (Tanggal|Nama|Keterangan) -> list dict."""
    ws = None
    for n in wb.sheetnames:
        if n.lower() == "catatan":
            ws = wb[n]; break
    if ws is None:
        return []
    out = []
    for r in range(2, ws.max_row + 1):
        tgl = ws.cell(row=r, column=1).value
        nama = ws.cell(row=r, column=2).value
        ket = ws.cell(row=r, column=3).value
        if not ket:
            continue
        iso = None
        if hasattr(tgl, "strftime"):
            iso = tgl.strftime("%Y-%m-%d")
        elif isinstance(tgl, str) and tgl.strip():
            iso = tgl.strip()
        out.append({
            "tanggal": iso,
            "nama": str(nama).strip() if nama else "",
            "keterangan": str(ket).strip(),
        })
    return out

if __name__ == "__main__":
    if not XLSM:
        print("Usage: sync-excel.py <file.xlsm>"); sys.exit(1)
    r = parse(XLSM)
    r2 = parse_catatan(openpyxl.load_workbook(XLSM, data_only=True, read_only=True))
    out = {
        "version": 1,
        "updated_at": datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7))).isoformat(timespec="seconds"),
        "bulan": r["bulan"],
        "attendance": r["attendance"],
        "catatan": r2,
    }
    json.dump(out, open(OUT, "w"), ensure_ascii=False)
    print(f"OK: {r['n_data']} data, {r['n_siswa']} siswa, {len(r['attendance'])} tanggal, bulan {r['bulan']}")
    print("->", OUT)
